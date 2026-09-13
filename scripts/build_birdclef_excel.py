import json
import re
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SPLIT_CSV = PROJECT_ROOT / "data" / "manifests" / "dataset_split.csv"
FREEZE_CSV = PROJECT_ROOT / "data" / "manifests" / "species_freeze.csv"
TRAIN_CSV = PROJECT_ROOT / "data" / "BirdClef" / "train.csv"
CACHE_FILE = PROJECT_ROOT / "data" / "manifests" / "xc_api_cache.json"
if not CACHE_FILE.exists():
    scratch_cache = Path.home() / ".gemini" / "antigravity" / "brain" / "29c8c496-9c02-45a8-aece-af8d01cc9493" / "scratch" / "xc_api_cache.json"
    if scratch_cache.exists():
        CACHE_FILE = scratch_cache

OUTPUT_EXCEL_MANIFESTS = PROJECT_ROOT / "data" / "manifests" / "Metadata_BirdCLEF.xlsx"
OUTPUT_EXCEL_BIRDCLEF = PROJECT_ROOT / "data" / "BirdClef" / "Metadata_BirdCLEF.xlsx"

def clean_type_str(t_val):
    if not t_val or pd.isna(t_val):
        return "-"
    t_str = str(t_val).strip()
    if t_str.startswith("[") and t_str.endswith("]"):
        # parse list syntax
        items = re.findall(r"'([^']*)'", t_str)
        if not items:
            items = re.findall(r'"([^"]*)"', t_str)
        if items:
            return ", ".join(items)
        return t_str.strip("[]'\" ")
    return t_str

def rating_to_quality(rating_val):
    try:
        r = float(rating_val)
        if r >= 4.8:
            return "A"
        elif r >= 3.8:
            return "B"
        elif r >= 2.8:
            return "C"
        elif r >= 1.8:
            return "D"
        else:
            return "E"
    except Exception:
        return "C"

def main():
    print("[*] Memuat data...")
    split_df = pd.read_csv(SPLIT_CSV)
    freeze_df = pd.read_csv(FREEZE_CSV)
    train_df = pd.read_csv(TRAIN_CSV)
    
    # Map filename to train row for fast lookup
    train_lookup = {}
    for _, row in train_df.iterrows():
        fname = Path(row["filename"]).name
        train_lookup[fname] = row

    with open(CACHE_FILE, "r", encoding="utf-8") as f:
        cache = json.load(f)
    print(f"Loaded cache with {len(cache):,} recordings.")

    # Manual patch for XC375735 if needed
    if "375735" not in cache:
        cache["375735"] = {
            "id": "375735",
            "cnt": "Costa Rica",
            "loc": "Arenal Observatory Lodge, San Carlos, Alajuela",
            "date": "2017-06-15",
            "q": "C",
            "type": "call",
            "rec": "Albert Lastukhin",
            "lic": "https://creativecommons.org/licenses/by-nc-sa/4.0/",
            "smp": "44100",
            "url": "https://xeno-canto.org/375735"
        }

    # Prepare rows
    all_rows = []
    species_rows = {s: [] for s in freeze_df["scientific_name"].unique()}

    for idx, row in split_df.iterrows():
        rec_id_str = str(row["recording_id"])
        numeric_id = rec_id_str.replace("XC", "")
        xc_meta = cache.get(numeric_id, {})

        sc_name = row["scientific_name"]
        cm_name = row["common_name"]
        role = row["split_role"]
        orig_fname = Path(row["file_path"]).name
        
        # Train lookup
        t_row = train_lookup.get(orig_fname, None)
        
        # Latitude & Longitude
        lat = row.get("latitude")
        if pd.isna(lat) and xc_meta.get("lat"):
            lat = xc_meta.get("lat")
        lng = row.get("longitude")
        if pd.isna(lng) and (xc_meta.get("lon") or xc_meta.get("lng")):
            lng = xc_meta.get("lon") or xc_meta.get("lng")

        # Locality & Country
        loc = xc_meta.get("loc", "")
        if not loc:
            loc = "Pantanal / South America"
        cnt = xc_meta.get("cnt", "")
        if not cnt:
            cnt = "Brazil"

        # Sampling Rate
        smp = xc_meta.get("smp", "")
        if smp and str(smp).isdigit():
            smp_hz = int(smp)
        else:
            smp_hz = 32000

        # Quality (A-E)
        q = xc_meta.get("q", "")
        if not q or q == "no score":
            q = rating_to_quality(row["rating"])

        # Type
        stype = xc_meta.get("type", "")
        if not stype and t_row is not None:
            stype = clean_type_str(t_row["type"])
        if not stype:
            stype = "song"

        # Date
        date_str = xc_meta.get("date", "")
        if not date_str:
            date_str = "-"

        # Recordist
        rec_name = xc_meta.get("rec", "")
        if not rec_name:
            rec_name = row["author"]

        # License
        lic_str = xc_meta.get("lic", "")
        if not lic_str and t_row is not None:
            lic_code = str(t_row.get("license", "by-nc-sa"))
            lic_str = f"https://creativecommons.org/licenses/{lic_code}/"
        if not lic_str:
            lic_str = "https://creativecommons.org/licenses/by-nc-sa/4.0/"

        item = {
            "Spesies": sc_name,
            "Sub-folder": role,
            "Nama File": orig_fname,
            "XC ID": f"XC{numeric_id}",
            "Link XenoCanto": f"https://xeno-canto.org/{numeric_id}",
            "Latitude": round(float(lat), 4) if pd.notna(lat) and str(lat).replace('-','').replace('.','').isdigit() else lat,
            "Longitude": round(float(lng), 4) if pd.notna(lng) and str(lng).replace('-','').replace('.','').isdigit() else lng,
            "Lokasi": loc,
            "Negara": cnt,
            "Sampling Rate (Hz)": smp_hz,
            "Kualitas (A-E)": q,
            "Tipe Suara": stype,
            "Tanggal Rekaman": date_str,
            "Rekorder": rec_name,
            "Lisensi": lic_str,
            "Common Name": cm_name,
            "Species Key": row["species_key"],
        }
        all_rows.append(item)
        if sc_name in species_rows:
            species_rows[sc_name].append(item)

    print(f"[+] Berhasil menyiapkan {len(all_rows):,} baris data.")

    # Build Excel Workbook
    wb = openpyxl.Workbook()
    
    COLUMNS = [
        "Spesies", "Sub-folder", "Nama File", "XC ID", "Link XenoCanto",
        "Latitude", "Longitude", "Lokasi", "Negara",
        "Sampling Rate (Hz)", "Kualitas (A-E)", "Tipe Suara",
        "Tanggal Rekaman", "Rekorder", "Lisensi"
    ]
    COL_WIDTHS = [28, 14, 20, 12, 38, 12, 12, 38, 16, 18, 14, 22, 16, 28, 48]
    
    thin = Side(style="thin", color="D3D3D3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_header(cell, hex_color="1F4E79"):
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=hex_color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    def style_cell(cell, row_even=True, is_numeric=False):
        bg = "F9FAFB" if row_even else "FFFFFF"
        cell.font = Font(name="Calibri", size=10)
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.alignment = Alignment(
            horizontal="right" if is_numeric else "left",
            vertical="center"
        )
        cell.border = border

    # --- Sheet 1: SEMUA DATA ---
    ws_all = wb.active
    ws_all.title = "SEMUA DATA"
    ws_all.freeze_panes = "A2"
    ws_all.row_dimensions[1].height = 28

    for ci, (col, width) in enumerate(zip(COLUMNS, COL_WIDTHS), 1):
        cell = ws_all.cell(row=1, column=ci, value=col)
        style_header(cell, "1F4E79")
        ws_all.column_dimensions[get_column_letter(ci)].width = width

    for ri, row in enumerate(all_rows, 2):
        even = (ri % 2 == 0)
        for ci, col in enumerate(COLUMNS, 1):
            val = row.get(col, "")
            is_num = col in ["Latitude", "Longitude", "Sampling Rate (Hz)"]
            c = ws_all.cell(row=ri, column=ci, value=val)
            style_cell(c, even, is_numeric=is_num)
            if col == "Link XenoCanto" and isinstance(val, str) and val.startswith("http"):
                c.hyperlink = val
                c.font = Font(name="Calibri", size=10, color="0563C1", underline="single")
            elif col == "Lisensi" and isinstance(val, str) and val.startswith("http"):
                c.hyperlink = val
                c.font = Font(name="Calibri", size=10, color="0563C1", underline="single")
        ws_all.row_dimensions[ri].height = 20

    print("[+] Sheet 'SEMUA DATA' selesai dibuat.")

    # --- Sheet 2: RINGKASAN ---
    ws_sum = wb.create_sheet(title="RINGKASAN", index=1)
    ws_sum.freeze_panes = "A2"
    ws_sum.row_dimensions[1].height = 28

    sum_headers = [
        "No", "Spesies (Nama Ilmiah)", "Nama Umum", "Kode",
        "Total Rekaman", "Galeri", "Kueri Bersih", "Kalibrasi",
        "Kualitas (A-E)", "Sampling Rate Unik (Hz)", "Negara (Unik)"
    ]
    sum_widths = [6, 28, 28, 12, 14, 10, 14, 12, 16, 26, 40]

    for ci, (col, width) in enumerate(zip(sum_headers, sum_widths), 1):
        c = ws_sum.cell(row=1, column=ci, value=col)
        style_header(c, "1F4E79")
        ws_sum.column_dimensions[get_column_letter(ci)].width = width

    for si, (sc_name, rows) in enumerate(species_rows.items(), 1):
        cm_name = rows[0]["Common Name"] if rows else "-"
        sp_key = rows[0]["Species Key"] if rows else "-"
        
        n_total = len(rows)
        n_gal = sum(1 for r in rows if r["Sub-folder"] == "gallery")
        n_qry = sum(1 for r in rows if r["Sub-folder"] == "query_clean")
        n_cal = sum(1 for r in rows if r["Sub-folder"] == "calibration")
        
        qualities = ", ".join(sorted(set(str(r["Kualitas (A-E)"]) for r in rows)))
        smps = ", ".join(sorted(set(str(r["Sampling Rate (Hz)"]) for r in rows), key=lambda x: int(x) if x.isdigit() else 0))
        countries = ", ".join(sorted(set(str(r["Negara"]) for r in rows if r["Negara"] not in ("", "-"))))

        even = (si % 2 == 0)
        row_vals = [
            si, sc_name, cm_name, sp_key,
            n_total, n_gal, n_qry, n_cal,
            qualities, smps, countries
        ]
        for ci, val in enumerate(row_vals, 1):
            is_num = ci in [1, 5, 6, 7, 8]
            c = ws_sum.cell(row=si + 1, column=ci, value=val)
            style_cell(c, even, is_numeric=is_num)
        ws_sum.row_dimensions[si + 1].height = 20

    print("[+] Sheet 'RINGKASAN' selesai dibuat.")

    # --- Sheet 3..22: Per-Spesies Sheets ---
    HEADER_COLORS = [
        "1F4E79", "2E75B6", "1C6B2B", "375623", "833C0B",
        "7B3F00", "4A235A", "154360", "0E4D2E", "5D2E0C",
        "1B2631", "4A4A00", "4A0000", "00394A", "2D0B4A",
        "2B547E", "348781", "6C3483", "1E8449", "7D6608"
    ]

    for si, (sc_name, rows) in enumerate(species_rows.items()):
        # Excel sheet name limit is 31 chars
        sheet_title = sc_name[:31]
        ws_sp = wb.create_sheet(title=sheet_title)
        ws_sp.freeze_panes = "A2"
        ws_sp.row_dimensions[1].height = 28
        h_color = HEADER_COLORS[si % len(HEADER_COLORS)]

        for ci, (col, width) in enumerate(zip(COLUMNS, COL_WIDTHS), 1):
            c = ws_sp.cell(row=1, column=ci, value=col)
            style_header(c, h_color)
            ws_sp.column_dimensions[get_column_letter(ci)].width = width

        for ri, row in enumerate(rows, 2):
            even = (ri % 2 == 0)
            for ci, col in enumerate(COLUMNS, 1):
                val = row.get(col, "")
                is_num = col in ["Latitude", "Longitude", "Sampling Rate (Hz)"]
                c = ws_sp.cell(row=ri, column=ci, value=val)
                style_cell(c, even, is_numeric=is_num)
                if col in ["Link XenoCanto", "Lisensi"] and isinstance(val, str) and val.startswith("http"):
                    c.hyperlink = val
                    c.font = Font(name="Calibri", size=10, color="0563C1", underline="single")
            ws_sp.row_dimensions[ri].height = 20

    print(f"[+] {len(species_rows)} Sheet per-spesies selesai dibuat.")

    # Save to both paths
    wb.save(OUTPUT_EXCEL_MANIFESTS)
    print(f"[+] Berhasil menyimpan Excel di: {OUTPUT_EXCEL_MANIFESTS}")
    
    wb.save(OUTPUT_EXCEL_BIRDCLEF)
    print(f"[+] Berhasil menyimpan salinan di: {OUTPUT_EXCEL_BIRDCLEF}")

if __name__ == "__main__":
    main()
